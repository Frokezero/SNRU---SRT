import openpyxl
import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def read_excel():
    path = r"D:\โฟลเดอร์ใหม่ (6)\รายชื่อนักศึกษาปี69.xlsx"
    if not os.path.exists(path):
        print("Excel file not found!")
        return
        
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active
        print(f"Sheet Name: {sheet.title}")
        print(f"Max Row: {sheet.max_row}")
        
        # Read unique majors in column 6 (1-based index)
        majors = set()
        for r in range(2, sheet.max_row + 1):
            val = sheet.cell(row=r, column=6).value
            if val:
                majors.add(str(val).strip())
                
        print("\nUnique majors in Excel column 6:")
        for m in sorted(majors):
            print(f"- {repr(m)}")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    read_excel()

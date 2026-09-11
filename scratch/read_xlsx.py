import openpyxl
import sys
import io

# Setup UTF-8 output encoding for Windows command line print
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook(r"D:\โฟลเดอร์ใหม่ (6)\รายชื่อนักศึกษาปี69.xlsx", data_only=True)
sheet = wb.active

print(f"Sheet Title: {sheet.title}")
print(f"Total Rows: {sheet.max_row}")

# Print headers
headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
print(f"Headers: {headers}")

# Print first 5 data rows
for r in range(2, 7):
    row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
    print(f"Row {r}: {row_vals}")

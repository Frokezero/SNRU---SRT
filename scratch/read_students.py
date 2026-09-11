import openpyxl

wb = openpyxl.load_workbook(r"D:\โฟลเดอร์ใหม่ (6)\รายชื่อนักศึกษาปี69.xlsx", data_only=True)
print(f"Sheets: {wb.sheetnames}")

sheet = wb.active
print(f"Active Sheet Title: {sheet.title}")

# Print first 15 rows
for r in range(1, 16):
    row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
    print(f"Row {r}: {row_vals}")

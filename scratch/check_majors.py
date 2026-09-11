import openpyxl
import os

filepath = r"D:\โฟลเดอร์ใหม่ (6)\รายชื่อนักศึกษาปี69.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)
sheet = wb.active

majors = set()
for r in range(2, sheet.max_row + 1):
    val = sheet.cell(row=r, column=6).value
    if val:
        majors.add(val)

# Configure output to support utf-8 print on Windows
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

print("Unique Majors in Excel:")
for m in sorted(majors):
    print(m)

print(f"Total students in Excel: {sheet.max_row - 1}")
